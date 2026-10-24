# -*- coding: utf-8 -*-
"""把 pptx 渲染成全页 HTML 近似预览，用于检查布局、溢出、对齐和图表数据表达。

用法:
  python preview_ppt.py <file.pptx> <out.html> [--slide N] [--accent 283E4E]
说明:
  默认输出全部页面；--slide 使用从 0 开始的页码。
  HTML 是近似渲染，不等于 PowerPoint 最终效果，复杂组合形状、嵌入对象和动画需人工复核。
"""
import argparse
import base64
import html as _h
import os
import sys

sys.stdout.reconfigure(encoding="utf-8")

try:
    from pptx import Presentation
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.dml import MSO_FILL
    from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
    from pptx.enum.text import MSO_ANCHOR
    from pptx.oxml.ns import qn
except ImportError:
    print("需要 python-pptx：uv run --python 3.14 --with python-pptx python preview_ppt.py")
    raise SystemExit(2)


EMU_IN = 914400.0
DPI = 96.0
PALETTE = ["283E4E", "B7C1C8", "000000", "F1F5F8"]


def normalize_hex(value, default):
    raw = str(value or default).strip().lstrip("#").upper()
    return raw if len(raw) == 6 and all(c in "0123456789ABCDEF" for c in raw) else default


def rgb(value, default="000000"):
    try:
        return normalize_hex(str(value.rgb), default)
    except Exception:
        return default


EA_MAP = {"宋体": "SimSun", "楷体": "KaiTi", "黑体": "SimHei"}


def ea_face(run):
    """读 a:ea（中文字形槽）；python-pptx 的 font.name 只管 a:latin（西文字形槽）。"""
    rPr = run._r.find(qn("a:rPr"))
    if rPr is None:
        return ""
    ea = rPr.find(qn("a:ea"))
    return ea.get("typeface") if ea is not None else ""


def font_stack(latin, ea):
    """按 latin + ea 双槽拼 CSS 字体栈，让预览与 PowerPoint 的实际渲染一致。"""
    parts = []
    if latin:
        parts.append('"%s"' % str(latin).replace('"', ""))
    if ea:
        parts.append('"%s"' % EA_MAP.get(ea, ea).replace('"', ""))
    return ",".join(parts) or '"Microsoft YaHei"'


def css_font(face, size_pt, color, bold=False, ea=""):
    return "font-family:%s;font-size:%.2fpx;color:#%s;font-weight:%s" % (
        font_stack(face, ea), max(float(size_pt or 16), 1) * DPI / 72,
        normalize_hex(color, "000000"), "bold" if bold else "normal"
    )


SHAPE_POLY = {
    MSO_SHAPE.ISOSCELES_TRIANGLE: "50,0 100,100 0,100",
    MSO_SHAPE.HEXAGON: "25,0 75,0 100,50 75,100 25,100 0,50",
    MSO_SHAPE.PARALLELOGRAM: "14,0 100,0 86,100 0,100",
    MSO_SHAPE.CHEVRON: "0,0 74,0 100,50 74,100 0,100 26,50",
    MSO_SHAPE.PENTAGON: "0,0 72,0 100,50 72,100 0,100",
    MSO_SHAPE.DIAMOND: "50,0 100,50 50,100 0,50",
    MSO_SHAPE.RIGHT_ARROW: "0,25 60,25 60,0 100,50 60,100 60,75 0,75",
    MSO_SHAPE.TRAPEZOID: "18,0 82,0 100,100 0,100",
}
SHAPE_CLIP = {
    MSO_SHAPE.OVAL: "border-radius:50%;",
    MSO_SHAPE.CAN: "border-radius:18% / 8%;",
    MSO_SHAPE.ROUNDED_RECTANGLE: "border-radius:10px;",
    MSO_SHAPE.FLOWCHART_PROCESS: "border-radius:10px;",
}


def poly_clip(points):
    """把 '0,0 74,0' 转成 clip-path 用的百分比串。"""
    pairs = []
    for item in points.split():
        x, y = item.split(",")
        pairs.append("%s%% %s%%" % (x, y))
    return "clip-path:polygon(%s);" % ",".join(pairs)


def chart_values(chart):
    try:
        categories = [str(c) for c in chart.plots[0].categories]
    except Exception:
        categories = []
    series = []
    try:
        for ser in chart.series:
            series.append([float(v) for v in ser.values])
    except Exception:
        pass
    return categories, series


def scatter_xy(chart):
    xs, ys = [], []
    try:
        ser = chart.series[0]
        xs = [float(v) for v in ser.x_values]
        ys = [float(v) for v in ser.y_values]
    except Exception:
        try:
            root = chart._chartSpace
            node = root.find(".//" + qn("c:ser"))
            x_cache = node.find(qn("c:xVal") + "/" + qn("c:numRef") + "/" + qn("c:numCache"))
            y_cache = node.find(qn("c:yVal") + "/" + qn("c:numRef") + "/" + qn("c:numCache"))
            for cache, target in ((x_cache, xs), (y_cache, ys)):
                if cache is not None:
                    for pt in cache.findall(qn("c:pt")):
                        value = pt.find(qn("c:v"))
                        if value is not None and value.text:
                            target.append(float(value.text))
        except Exception:
            pass
    return xs, ys


def chart_html(chart, width, height, accent, band):
    chart_type = chart.chart_type
    categories, series = chart_values(chart)
    values = series[0] if series else []
    esc = _h.escape
    body = []

    if chart_type in (XL_CHART_TYPE.DOUGHNUT, XL_CHART_TYPE.PIE):
        vals = [abs(float(v or 0)) for v in values]
        total = sum(vals) or 1.0
        start = 0.0
        stops = []
        for i, value in enumerate(vals):
            end = start + value / total * 100.0
            color = PALETTE[i % len(PALETTE)]
            stops.append("#%s %.2f%% %.2f%%" % (color, start, end))
            start = end
        size = min(width, height) * 0.72
        body.append(
            "<div style='width:%.1fpx;height:%.1fpx;border-radius:50%%;background:conic-gradient(%s);"
            "margin:auto;position:relative'><div style='position:absolute;inset:22%%;border-radius:50%%;"
            "background:#fff'></div></div>" % (size, size, ",".join(stops))
        )
    elif "SCATTER" in str(chart_type):
        xs, ys = scatter_xy(chart)
        if xs and ys:
            x0, x1 = min(xs), max(xs)
            y0, y1 = min(ys), max(ys)
            dx = (x1 - x0) or 1.0
            dy = (y1 - y0) or 1.0
            for x, y in zip(xs, ys):
                px = 34 + (x - x0) / dx * max(width - 68, 1)
                py = height - 28 - (y - y0) / dy * max(height - 58, 1)
                body.append(
                    "<div style='position:absolute;left:%.1fpx;top:%.1fpx;width:8px;height:8px;"
                    "border-radius:50%%;background:#%s'></div>" % (px - 4, py - 4, accent)
                )
        else:
            body.append("<div style='padding:20px'>散点图数据无法近似预览</div>")
    elif chart_type in (XL_CHART_TYPE.LINE, XL_CHART_TYPE.LINE_MARKERS) and values:
        y0 = min([0.0] + values)
        y1 = max([0.0] + values)
        dy = (y1 - y0) or 1.0
        points = []
        for i, value in enumerate(values):
            x = 34 + (i / max(len(values) - 1, 1)) * max(width - 68, 1)
            y = height - 28 - (value - y0) / dy * max(height - 58, 1)
            points.append("%.1f,%.1f" % (x, y))
        body.append(
            "<svg width='%.1f' height='%.1f'><polyline points='%s' fill='none' stroke='#%s' stroke-width='3'/>"
            "</svg>" % (width, height, " ".join(points), accent)
        )
    elif chart_type == XL_CHART_TYPE.COLUMN_CLUSTERED:
        vals = [float(v or 0) for v in values]
        data_min = min([0.0] + vals)
        data_max = max([0.0] + vals)
        span = (data_max - data_min) or 1.0
        top_pad, bottom_pad = 28, 34
        avail = max(height - top_pad - bottom_pad, 1)
        zero_y = top_pad + (data_max / span) * avail
        count = max(len(vals), 1)
        bar_w = width / count * 0.52
        gap = (width - count * bar_w) / (count + 1)
        for i, (category, value) in enumerate(zip(categories, vals)):
            x = gap + i * (bar_w + gap)
            bar_h = abs(value) / span * avail
            y = zero_y - bar_h if value >= 0 else zero_y
            body.append(
                "<div style='position:absolute;left:%.1fpx;top:%.1fpx;width:%.1fpx;height:%.1fpx;background:#%s'></div>"
                % (x, y, bar_w, max(bar_h, 1), accent)
            )
            body.append(
                "<div style='position:absolute;left:%.1fpx;top:%.1fpx;width:%.1fpx;text-align:center;%s'>%s</div>"
                % (x - 8, y - 22, bar_w + 16, css_font("", 16, accent, True), _h.escape("%g" % value))
            )
            body.append(
                "<div style='position:absolute;left:%.1fpx;top:%.1fpx;width:%.1fpx;text-align:center;%s'>%s</div>"
                % (x - 10, height - 24, bar_w + 20, css_font("", 14, "000000"), esc(category))
            )
    else:
        vals = [float(v or 0) for v in values]
        data_min = min([0.0] + vals)
        data_max = max([0.0] + vals)
        span = (data_max - data_min) or 1.0
        label_w = min(width * 0.34, 180)
        plot_x = label_w + 10
        plot_w = max(width - plot_x - 16, 1)
        zero_x = plot_x + ((0 - data_min) / span) * plot_w
        row_h = height / max(len(vals), 1)
        for i, (category, value) in enumerate(zip(categories, vals)):
            y = i * row_h + row_h * 0.25
            bar_h = min(row_h * 0.5, 38)
            value_x = plot_x + ((min(value, 0) - data_min) / span) * plot_w
            bar_w = abs(value) / span * plot_w
            body.append(
                "<div style='position:absolute;left:%.1fpx;top:%.1fpx;width:%.1fpx;text-align:right;%s'>%s</div>"
                % (0, y + 4, label_w, css_font("", 14, "000000"), esc(category))
            )
            body.append(
                "<div style='position:absolute;left:%.1fpx;top:%.1fpx;width:%.1fpx;height:%.1fpx;background:#%s'></div>"
                % (value_x, y, max(bar_w, 1), bar_h, accent)
            )
            body.append(
                "<div style='position:absolute;left:%.1fpx;top:%.1fpx;%s'>%s</div>"
                % (max(value_x + bar_w + 6, zero_x + 6), y + 4, css_font("", 14, accent, True), "%g" % value)
            )

    return "<div class='chart'>%s</div>" % "".join(body)


def image_data(image):
    content_type = getattr(image, "content_type", "") or "image/jpeg"
    extension = (getattr(image, "ext", "") or "").lower()
    if extension in ("png", "gif", "bmp", "webp"):
        content_type = "image/" + extension
    return content_type, base64.b64encode(image.blob).decode("ascii")


def text_html(text_frame, font_default=""):
    anchor = getattr(text_frame, "vertical_anchor", None)
    justify = {MSO_ANCHOR.TOP: "flex-start", MSO_ANCHOR.MIDDLE: "center", MSO_ANCHOR.BOTTOM: "flex-end"}.get(
        anchor, "flex-start"
    )
    paragraphs = []
    for paragraph in text_frame.paragraphs:
        runs = []
        for run in paragraph.runs:
            face = run.font.name or font_default
            ea = ea_face(run)
            size = run.font.size.pt if run.font.size else 16
            color = rgb(run.font.color, "000000")
            bold = bool(run.font.bold)
            runs.append(
                "<span style='%s'>%s</span>" % (css_font(face, size, color, bold, ea), _h.escape(run.text))
            )
        alignment = {1: "left", 2: "center", 3: "right"}.get(int(paragraph.alignment or 1), "left")
        paragraphs.append("<div style='width:100%%;text-align:%s;line-height:1.25'>%s</div>" % (alignment, "".join(runs)))
    return "<div class='text'>%s</div>" % "".join(paragraphs)


def render_table(shape, width, height):
    table = shape.table
    column_total = sum(column.width for column in table.columns) or 1
    cols = "".join("<col style='width:%.3f%%'>" % (column.width / column_total * 100) for column in table.columns)
    rows = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            background = "FFFFFF"
            try:
                if cell.fill.type == MSO_FILL.SOLID:
                    background = rgb(cell.fill.fore_color, "FFFFFF")
            except Exception:
                pass
            paragraphs = []
            for paragraph in cell.text_frame.paragraphs:
                runs = []
                for run in paragraph.runs:
                    face = run.font.name or ""
                    ea = ea_face(run)
                    size = run.font.size.pt if run.font.size else 16
                    color = rgb(run.font.color, "000000")
                    runs.append(
                        "<span style='%s'>%s</span>"
                        % (css_font(face, size, color, bool(run.font.bold), ea), _h.escape(run.text))
                    )
                alignment = {1: "left", 2: "center", 3: "right"}.get(int(paragraph.alignment or 1), "left")
                paragraphs.append("<div style='text-align:%s'>%s</div>" % (alignment, "".join(runs)))
            cells.append(
                "<td style='background:#%s;padding:6px 8px;vertical-align:middle;border:1px solid #D5DCE1'>%s</td>"
                % (background, "".join(paragraphs))
            )
        rows.append("<tr>%s</tr>" % "".join(cells))
    return (
        "<div class='chart'><table style='width:100%%;height:%.1fpx;table-layout:fixed;border-collapse:collapse'>"
        "<colgroup>%s</colgroup>%s</table></div>" % (height, cols, "".join(rows))
    )


def render_slide(prs, slide, slide_number, accent, band):
    width = prs.slide_width / EMU_IN * DPI
    height = prs.slide_height / EMU_IN * DPI
    parts = [
        "<section class='slide-wrap'><div class='slide-label'>第 %d 页</div>" % slide_number,
        "<div class='slide' style='width:%.2fpx;height:%.2fpx'>" % (width, height),
    ]

    for shape in slide.shapes:
        left = shape.left / EMU_IN * DPI
        top = shape.top / EMU_IN * DPI
        shape_w = shape.width / EMU_IN * DPI
        shape_h = shape.height / EMU_IN * DPI
        style = "left:%.2fpx;top:%.2fpx;width:%.2fpx;height:%.2fpx;" % (left, top, shape_w, shape_h)

        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            parts.append(
                "<div class='shape warning' style='%s'>组合形状：请在 PowerPoint 中复核内部元素</div>" % style
            )
            continue
        if getattr(shape, "has_table", False):
            parts.append("<div class='shape' style='%s'>%s</div>" % (style, render_table(shape, shape_w, shape_h)))
            continue
        if getattr(shape, "has_chart", False):
            parts.append("<div class='shape' style='%s'>%s</div>" % (style, chart_html(shape.chart, shape_w, shape_h, accent, band)))
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            content_type, encoded = image_data(shape.image)
            image_w, image_h = shape.image.size
            crop_l = float(getattr(shape, "crop_left", 0) or 0)
            crop_t = float(getattr(shape, "crop_top", 0) or 0)
            crop_r = float(getattr(shape, "crop_right", 0) or 0)
            crop_b = float(getattr(shape, "crop_bottom", 0) or 0)
            visible_w = max(image_w * (1 - crop_l - crop_r), 1)
            visible_h = max(image_h * (1 - crop_t - crop_b), 1)
            scale = max(shape_w / visible_w, shape_h / visible_h)
            parts.append(
                "<div class='shape image' style='%s'><img src='data:%s;base64,%s' "
                "style='position:absolute;left:%.2fpx;top:%.2fpx;width:%.2fpx;height:%.2fpx'></div>"
                % (style, content_type, encoded, -crop_l * image_w * scale, -crop_t * image_h * scale,
                   image_w * scale, image_h * scale)
            )
            continue

        auto_shape = None
        try:
            auto_shape = shape.auto_shape_type
        except Exception:
            auto_shape = None
        polygon = SHAPE_POLY.get(auto_shape) if auto_shape is not None else None
        fill_color = None
        try:
            if shape.fill.type == MSO_FILL.SOLID:
                fill_color = rgb(shape.fill.fore_color, "FFFFFF")
        except Exception:
            pass
        border_color = None
        line_width = 1.0
        try:
            if shape.line.fill.type == MSO_FILL.SOLID:
                border_color = rgb(shape.line.color, "B7C1C8")
                line_width = shape.line.width.pt if shape.line.width else 1.0
        except Exception:
            pass
        clip = SHAPE_CLIP.get(auto_shape, "") if auto_shape is not None else ""
        rotation = getattr(shape, "rotation", 0) or 0
        transform = "transform:rotate(%.1fdeg);transform-origin:50%% 50%%;" % rotation if rotation else ""
        content = text_html(shape.text_frame) if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip() else ""
        if polygon and border_color and not content:
            # 多边形 + 描边用 SVG，否则 CSS clip-path 会把边框一起裁掉
            parts.append("<div class='shape' style='%s%s'><svg width='100%%' height='100%%' viewBox='0 0 100 100' "
                         "preserveAspectRatio='none'><polygon points='%s' fill='#%s' stroke='#%s' stroke-width='%.2f' "
                         "vector-effect='non-scaling-stroke'/></svg></div>"
                         % (style, transform, polygon, fill_color or "FFFFFF", border_color, max(line_width, 0.75)))
            continue
        if polygon:
            clip = poly_clip(polygon)
        fill = "background:#%s;" % fill_color if fill_color else ""
        border = "border:%.2fpx solid #%s;" % (max(line_width, 0.75), border_color) if border_color else ""
        parts.append("<div class='shape' style='%s%s%s%s%s'>%s</div>" % (style, fill, border, clip, transform, content))

    parts.append("</div></section>")
    return "".join(parts)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("src")
    parser.add_argument("out")
    parser.add_argument("--slide", type=int, default=None, help="从 0 开始的页码；省略时输出全部页面")
    parser.add_argument("--accent", default="283E4E")
    parser.add_argument("--band", default="F1F5F8")
    args = parser.parse_args()

    accent = normalize_hex(args.accent, "283E4E")
    band = normalize_hex(args.band, "F1F5F8")
    prs = Presentation(args.src)
    if not len(prs.slides):
        raise SystemExit("PPTX 中没有页面")

    if args.slide is None:
        indexes = range(len(prs.slides))
    else:
        if args.slide < 0 or args.slide >= len(prs.slides):
            raise SystemExit("--slide 超出范围：%d" % args.slide)
        indexes = [args.slide]

    slide_html = "".join(render_slide(prs, prs.slides[i], i + 1, accent, band) for i in indexes)
    document = """<!DOCTYPE html>
<html lang="zh-CN"><head><meta charset="utf-8"><title>PPT 近似预览</title>
<style>
*{box-sizing:border-box}body{margin:0;padding:24px;background:#E9EDF0;font-family:"Microsoft YaHei",sans-serif}
.slide-wrap{margin:0 auto 28px;width:max-content}.slide-label{margin:0 0 8px;color:#283E4E;font-size:14px}
.slide{position:relative;overflow:hidden;background:#fff;box-shadow:0 6px 24px rgba(30,50,64,.18)}
.shape{position:absolute;overflow:hidden}.shape.image{background:#F1F5F8}.shape.warning{display:flex;align-items:center;justify-content:center;padding:12px;border:1px dashed #B7C1C8;color:#283E4E;font-size:14px;text-align:center}
.text{position:absolute;inset:0;display:flex;flex-direction:column;justify-content:flex-start;white-space:pre-wrap}
.chart{position:absolute;inset:0;overflow:hidden}.chart span{white-space:pre-wrap}
</style></head><body>
<div style='max-width:900px;margin:0 auto 24px;padding:12px 16px;background:#fff;border-left:4px solid #283E4E'>
这是布局近似预览，不等于 PowerPoint 最终渲染。字体、图表、组合形状、嵌入对象、动画和导出效果仍需在 PowerPoint 中复核。
</div>
%s</body></html>""" % slide_html

    out_dir = os.path.dirname(os.path.abspath(args.out))
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    with open(args.out, "w", encoding="utf-8") as handle:
        handle.write(document)
    print("html", os.path.abspath(args.out))


if __name__ == "__main__":
    main()
