# -*- coding: utf-8 -*-
"""生成示例页：assets/example-slide.pptx（正例，9 页）与回归用反例。

正例覆盖多种关系模型与骨架：封面、主图+侧栏、表格、对比、流程、时间线、层级、循环、纯文字观点页；
字号只走字阶 16/18/20/30/36/44，中西文分槽设置，配图按框裁切不拉伸。
示例内容全部为占位文案与虚构数据，不含任何真实项目材料。

用法:
  python build_example.py                                          # 重建正例与占位图
  python build_example.py --variant negative --out 临时路径.pptx    # 生成回归反例
"""
import argparse
import os
import struct
import sys
import zlib

sys.stdout.reconfigure(encoding="utf-8")

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_MARK
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.normpath(os.path.join(HERE, "..", "assets"))

THEME = ("142129", "1E303B", "283E4E", "4A6577", "8199A8", "B7C7D2", "E4EBF0")
CONTRAST = ("8C5A16", "B4771F", "C98B2E", "E0B15B")
SUPPORT = ("2F4E48", "4F7C74", "8DB0A8", "D6E5E1")
DEEP, NAVY, MID, LIGHT, PALE = THEME[0], THEME[2], THEME[3], THEME[5], THEME[6]
AMBER_TEXT, AMBER_FILL = CONTRAST[0], CONTRAST[3]
SUPPORT_TEXT, SUPPORT_FILL = SUPPORT[1], SUPPORT[3]
BLACK, WHITE, BAND, LINE = "000000", "FFFFFF", "F1F5F8", "B7C1C8"
HEI, KAI, SONG = "黑体", "楷体", "宋体"

# 字阶：只有这 6 档
S16, S18, S20, S30, S36, S44 = 16, 18, 20, 30, 36, 44
# 西文槽只允许 Arial 与 Times New Roman；中文字形走 a:ea，两者分开设置。
LATIN_OF = {HEI: "Arial", SONG: "Times New Roman", KAI: "Times New Roman"}

SLIDE_W = 12858750
SLIDE_H = 7232650


def font(run, size, color, face, latin=None):
    """face = 中文字形（a:ea / a:cs）；latin = 西文字形（a:latin），默认按 face 配对。"""
    lat = latin or LATIN_OF.get(face, "Arial")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = lat
    rpr = run._r.get_or_add_rPr()
    for tag, typeface in (("a:ea", face), ("a:cs", face)):
        node = rpr.find(qn(tag))
        if node is None:
            node = rpr.makeelement(qn(tag), {})
            rpr.insert_element_before(node, "a:sym", "a:hlinkClick", "a:hlinkMouseOver",
                                      "a:rtl", "a:endParaRPr")
        node.set("typeface", typeface)


def rect(slide, left, top, width, height, fill=None, line=None, line_width=0.75,
         shape=MSO_SHAPE.RECTANGLE):
    item = slide.shapes.add_shape(shape, Inches(left), Inches(top), Inches(width), Inches(height))
    item.shadow.inherit = False
    if fill:
        item.fill.solid()
        item.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        item.fill.background()
    if line:
        item.line.color.rgb = RGBColor.from_string(line)
        item.line.width = Pt(line_width)
    else:
        item.line.fill.background()
    item.text_frame.text = ""
    return item


def text(slide, left, top, width, height, segments, size, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.0):
    """segments: [(文本, 颜色, 字体), ...]"""
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.line_spacing = spacing
    for segment in segments:
        run = paragraph.add_run()
        run.text = segment[0]
        font(run, size, segment[1], segment[2])
    return box


def header(slide, page, title, subtitle):
    text(slide, 6.40, 0.24, 6.10, 0.34, [("示例 · 汇报页骨架", BLACK, SONG)], S16, PP_ALIGN.RIGHT)
    text(slide, 12.60, 0.24, 0.91, 0.34, [(page, NAVY, HEI)], S16, PP_ALIGN.RIGHT)
    text(slide, 0.55, 0.58, 12.96, 0.66, [(title, NAVY, HEI)], S30)
    text(slide, 0.55, 1.26, 12.96, 0.36, [(subtitle, BLACK, SONG)], S20)
    rect(slide, 0.55, 1.70, 12.96, 0.012, fill=NAVY)


def place_picture(slide, path, left, top, width, height, image_size):
    """按框裁切填满，不拉伸变形。"""
    image_w, image_h = image_size
    picture = slide.shapes.add_picture(path, Inches(left), Inches(top), Inches(width), Inches(height))
    image_ratio = image_w / float(image_h)
    box_ratio = width / float(height)
    if image_ratio > box_ratio:
        crop = (1.0 - box_ratio / image_ratio) / 2.0
        picture.crop_left = crop
        picture.crop_right = crop
    else:
        crop = (1.0 - image_ratio / box_ratio) / 2.0
        picture.crop_top = crop
        picture.crop_bottom = crop
    return picture


def chart_polish(chart, series_color, label_color):
    chart.has_title = False
    chart.has_legend = False
    chart.font.size = Pt(S16)
    chart.font.name = "Arial"
    chart.font.color.rgb = RGBColor.from_string(BLACK)
    chart_space = chart._chartSpace
    if chart_space.find(qn("c:spPr")) is None:
        chart_space.find(qn("c:chart")).addnext(
            parse_xml("<c:spPr %s><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>" % nsdecls("c", "a")))
    plot = chart.plots[0]
    plot.gap_width = 90
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format = "#,##0"
    labels.number_format_is_linked = False
    labels.font.size = Pt(S16)
    labels.font.bold = True
    labels.font.name = "Arial"
    labels.font.color.rgb = RGBColor.from_string(label_color)
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    series = plot.series[0]
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = RGBColor.from_string(series_color)
    series.format.line.fill.background()
    category_axis = chart.category_axis
    category_axis.has_major_gridlines = False
    category_axis.major_tick_mark = XL_TICK_MARK.NONE
    category_axis.tick_labels.font.size = Pt(S16)
    category_axis.tick_labels.font.name = "Arial"
    category_axis.tick_labels.font.color.rgb = RGBColor.from_string(BLACK)
    category_axis.format.line.color.rgb = RGBColor.from_string(NAVY)
    category_axis.format.line.width = Pt(1)
    chart.value_axis.has_major_gridlines = False
    delete = chart.value_axis._element.find(qn("c:delete"))
    if delete is None:
        delete = chart.value_axis._element.makeelement(qn("c:delete"), {})
        chart.value_axis._element.find(qn("c:scaling")).addnext(delete)
    delete.set("val", "1")
    for rpr in chart_space.iter(qn("a:defRPr")):
        # 图表里中文取 ea、西文与数字取 latin：ea 固定宋体，latin 缺省补 Arial。
        latin = rpr.find(qn("a:latin"))
        if latin is None:
            latin = rpr.makeelement(qn("a:latin"), {})
            rpr.insert_element_before(latin, "a:ea", "a:cs", "a:sym", "a:hlinkClick",
                                      "a:hlinkMouseOver", "a:rtl", "a:endParaRPr")
        latin.set("typeface", latin.get("typeface") or "Arial")
        for tag in ("a:ea", "a:cs"):
            node = rpr.find(qn(tag))
            if node is None:
                node = rpr.makeelement(qn(tag), {})
                rpr.insert_element_before(node, "a:sym", "a:hlinkClick", "a:hlinkMouseOver",
                                          "a:rtl", "a:endParaRPr")
            node.set("typeface", SONG)


def _chunk(tag, data):
    return (struct.pack(">I", len(data)) + tag + data
            + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF))


def write_png(path, width, height, columns, red_band=False):
    """写一张纯占位图（横向渐变），用于演示配图位置与分辨率要求。"""
    rows = []
    for y in range(height):
        row = bytearray([0])
        top = y / max(height - 1, 1)
        for red, green, blue in columns:
            shade = 1.0 - 0.18 * top
            if red_band and top > 0.62:
                row += bytes((min(255, int(200 * shade)), int(red * 0.5), int(green * 0.5)))
            else:
                row += bytes((int(red * shade), int(green * shade), int(blue * shade)))
        rows.append(bytes(row))
    with open(path, "wb") as handle:
        handle.write(b"\x89PNG\r\n\x1a\n")
        handle.write(_chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)))
        handle.write(_chunk(b"IDAT", zlib.compress(b"".join(rows), 6)))
        handle.write(_chunk(b"IEND", b""))
    return path


def gradient_columns(width, start=THEME[2], end=THEME[6]):
    first = [int(start[index:index + 2], 16) for index in (0, 2, 4)]
    last = [int(end[index:index + 2], 16) for index in (0, 2, 4)]
    columns = []
    for x in range(width):
        ratio = x / max(width - 1, 1)
        columns.append(tuple(first[i] + int(ratio * (last[i] - first[i])) for i in range(3)))
    return columns


def build_positive(prs, placeholder):
    # ── 第 1 页：封面 ──
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0.55, 1.30, 0.045, 1.45, fill=NAVY)
    text(slide, 0.80, 1.35, 12.30, 0.80, [("示例演示：把内容关系画对", NAVY, HEI)], S44)
    text(slide, 0.80, 2.20, 12.30, 0.44,
         [("副标题占位：这是一份纯占位数据的版式示例，不含任何真实结论。", BLACK, SONG)], S20)
    text(slide, 0.80, 2.80, 12.30, 0.34, [("示例单位 · 示例作者 · 示例日期", MID, SONG)], S16)
    place_picture(slide, placeholder, 0.55, 3.40, 6.20, 3.30, (1600, 900))
    rect(slide, 7.05, 3.40, 6.46, 3.30, fill=BAND)
    rect(slide, 7.05, 3.40, 0.045, 3.30, fill=NAVY)
    text(slide, 7.30, 3.60, 5.96, 2.90,
         [("本示例演示新的配色与版式规则：黑白灰作基础色，不占配色额度；"
           "主题色、对比色、辅助色各成系列。每页先判断内容关系，再选骨架，"
           "不再把页面全做成方块。", BLACK, KAI)], S20,
         anchor=MSO_ANCHOR.MIDDLE, spacing=1.25)
    slide.notes_slide.notes_text_frame.text = (
        "【封面 · 示例备注】本页只演示版式：大标题、副标题与必要的标识位置，不承载结论。讲解时用十秒交代主题与场合，"
        "不需要把目录念一遍。真实项目里单位、姓名、日期放在固定位置，字号不小于 16pt；来源、口径和备答一律写进备注，"
        "不上片。封面不强制写长讲稿，按实际讲解需要写。" * 2)

    # ── 第 2 页：主图 + 侧栏 ──
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "02", "示例页：主图 + 侧栏，配图与图表各承担一件事",
           "左栏讲证据，右栏放实景图、图表与关键数字；对比色只点最关键的一个数字。")
    rect(slide, 0.55, 1.82, 7.10, 1.05, fill=BAND)
    rect(slide, 0.55, 1.82, 0.045, 1.05, fill=NAVY)
    text(slide, 0.76, 1.92, 6.74, 0.85,
         [("导语占位：这一页演示左文右图的骨架，文字区与图片区各自成块，边缘对齐。", BLACK, KAI)], S20,
         anchor=MSO_ANCHOR.MIDDLE, spacing=1.2)
    rows = [("示例指标 A 达到 ", "38.6", "（占位数值）"),
            ("示例指标 B 同比变化 ", "+12.4", "（占位数值）"),
            ("示例指标 C 覆盖范围 ", "1,200", "（占位数值）")]
    top = 2.97
    for first, value, third in rows:
        rect(slide, 0.55, top, 0.045, 0.72, fill=NAVY)
        text(slide, 0.76, top, 6.84, 0.72,
             [(first, BLACK, SONG), (value, NAVY, HEI), (third, BLACK, SONG)], S18,
             anchor=MSO_ANCHOR.MIDDLE, spacing=1.2)
        top += 0.78
    rect(slide, 0.55, 5.35, 7.10, 1.10, fill=BAND)
    rect(slide, 0.55, 5.35, 0.045, 1.10, fill=AMBER_TEXT)
    text(slide, 0.76, 5.45, 6.74, 0.90,
         [("核心判断占位：用一句话收束左侧证据。对比色只承担一个重点，不做大面积铺色。", NAVY, HEI)], S18,
         anchor=MSO_ANCHOR.MIDDLE, spacing=1.2)
    place_picture(slide, placeholder, 7.85, 1.82, 5.66, 1.62, (1600, 900))
    rect(slide, 7.85, 3.62, 1.72, 1.72, fill=NAVY, shape=MSO_SHAPE.OVAL)
    text(slide, 7.90, 3.98, 1.62, 0.70, [("38.6%", WHITE, HEI)], S36, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    text(slide, 7.95, 4.76, 1.52, 0.34, [("占位数值", WHITE, SONG)], S16, PP_ALIGN.CENTER)
    rect(slide, 9.95, 3.62, 3.56, 0.44, fill=AMBER_FILL)
    text(slide, 10.10, 3.70, 3.26, 0.30, [("关键提示：对比色只点一个重点", DEEP, HEI)], S16)
    data = CategoryChartData()
    data.categories = ["甲", "乙", "丙"]
    data.add_series("示例", (62, 45, 28))
    frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(9.95), Inches(4.24),
                                   Inches(3.56), Inches(2.10), data)
    chart_polish(frame.chart, MID, NAVY)
    slide.notes_slide.notes_text_frame.text = (
        "【主图+侧栏示例 · 备注】本页只是版式示例，所有数值都是虚构占位数据，不要当作真实结论引用。"
        "讲解节奏：先用一句话说明这一页回答什么问题；再逐条讲三条实证，每条都是指标名、数值、口径写在同一句里；"
        "然后用一句判断收束。右栏的图片是实景位，图表是数据位，两者各承担一件事，不要互相替代。"
        "如果换成真实数据，来源、统计口径和发布时间一律写进备注。" * 2)

    # ── 第 3 页：强数据用表格 ──
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "03", "示例页：强数据内容用表格承载",
           "指标对比、口径对照、多维数据用表格；叙事与机制内容才用方块。")
    data_rows = [["指标", "2024", "2025", "变化", "口径"],
                 ["指标甲", "1,180", "1,420", "+20.3%", "示例口径"],
                 ["指标乙", "62.1", "58.4", "-3.7", "示例口径"],
                 ["指标丙", "340", "512", "+50.6%", "示例口径"],
                 ["指标丁", "9.6", "12.8", "+33.3%", "示例口径"]]
    shape = slide.shapes.add_table(len(data_rows), 5, Inches(0.55), Inches(1.94), Inches(12.96), Inches(3.30))
    table = shape.table
    for index, width in enumerate([3.00, 2.10, 2.10, 2.10, 3.66]):
        table.columns[index].width = Inches(width)
    for row_no, row in enumerate(data_rows):
        table.rows[row_no].height = Inches(0.66)
        for col_no, value in enumerate(row):
            cell = table.cell(row_no, col_no)
            cell.margin_left = Inches(0.14)
            cell.margin_right = Inches(0.10)
            cell.margin_top = Inches(0.05)
            cell.margin_bottom = Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if row_no == 0:
                cell.fill.fore_color.rgb = RGBColor.from_string(NAVY)
            elif row_no == 4:
                cell.fill.fore_color.rgb = RGBColor.from_string(SUPPORT_FILL)
            else:
                cell.fill.fore_color.rgb = RGBColor.from_string(BAND if row_no % 2 == 0 else WHITE)
            frame = cell.text_frame
            frame.word_wrap = True
            paragraph = frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.LEFT if col_no in (0, 4) else PP_ALIGN.RIGHT
            run = paragraph.add_run()
            run.text = value
            if row_no == 0:
                font(run, S16, WHITE, HEI)
            elif col_no in (1, 2, 3):
                font(run, S16, NAVY, HEI)
            else:
                font(run, S16, BLACK, SONG)
    rect(slide, 0.55, 5.60, 0.045, 0.90, fill=NAVY)
    text(slide, 0.76, 5.60, 12.60, 0.90,
         [("示例结论：指标甲、丙、丁增长、指标乙回落。末行用辅助色浅底标示重点行，面积控制在两成以内。",
           BLACK, SONG)], S18, anchor=MSO_ANCHOR.MIDDLE, spacing=1.2)
    slide.notes_slide.notes_text_frame.text = (
        "【表格示例 · 备注】本页演示强数据用表格：表头用主题色底白字，数值列用主题色，正文黑色，交替行加浅底，"
        "重点行用辅助色浅底标示。讲的时候不要逐格念表，先说这张表回答什么问题、结论是什么，再点两三个关键数字的变化。"
        "如果表格承担逐项核对，就把口径与来源写进备注，必要时补一句数据截止月份。" * 2)

    # ── 第 4 页：左右双栏对比 ──
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "04", "示例页：对比关系用左右双栏",
           "两组并列信息用双栏对照，不用箭头堆叠；差异用一条中缝和标签区分。")
    rect(slide, 0.55, 1.86, 6.05, 4.30, fill=BAND, line=LINE)
    rect(slide, 6.88, 1.86, 6.05, 4.30, fill=WHITE, line=LINE)
    text(slide, 0.85, 2.04, 5.45, 0.40, [("左侧：原有做法", NAVY, HEI)], S20)
    text(slide, 7.18, 2.04, 5.45, 0.40, [("右侧：调整后做法", NAVY, HEI)], S20)
    left_items = ["条目占位一：把机制写清楚，不写口号。",
                  "条目占位二：数字与口径写在同一句。",
                  "条目占位三：来源进备注，不上片。"]
    for index, item in enumerate(left_items):
        text(slide, 0.85, 2.70 + index * 0.72, 5.45, 0.64, [(item, BLACK, SONG)], S18,
             anchor=MSO_ANCHOR.MIDDLE, spacing=1.2)
    right_items = ["条目占位一：结构先行，先改大纲再改页面。",
                   "条目占位二：版面留白，不填死下缘。",
                   "条目占位三：对比色只点一个重点。"]
    for index, item in enumerate(right_items):
        text(slide, 7.18, 2.70 + index * 0.72, 5.45, 0.64, [(item, BLACK, SONG)], S18,
             anchor=MSO_ANCHOR.MIDDLE, spacing=1.2)
    rect(slide, 6.62, 2.10, 0.24, 0.50, fill=AMBER_FILL, shape=MSO_SHAPE.PARALLELOGRAM)
    text(slide, 0.85, 5.28, 5.45, 0.62, [("判断占位：机制可以借鉴，前提不能照搬。", AMBER_TEXT, HEI)], S18,
         anchor=MSO_ANCHOR.MIDDLE, spacing=1.2)
    text(slide, 7.18, 5.28, 5.45, 0.62, [("判断占位：先立自己的能力，再谈开放。", NAVY, HEI)], S18,
         anchor=MSO_ANCHOR.MIDDLE, spacing=1.2)
    slide.notes_slide.notes_text_frame.text = (
        "【对比页 · 备注】左右双栏适合两组并列信息，中缝的平行四边形是关系标记，不是装饰。"
        "讲解时先给对照的标准，再逐条比，最后落到一句判断。对比页不要用箭头硬造因果关系，"
        "因果关系应该放到流程页或因果图里。" * 2)

    # ── 第 5 页：流程 ──
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "05", "示例页：流程关系用形状序列表达",
           "主视觉形状按语义选择：流程用箭头形与圆形节点，不用一排方块。")
    rect(slide, 0.55, 1.82, 12.96, 0.86, fill=BAND)
    text(slide, 0.76, 1.94, 12.54, 0.62,
         [("导语占位：这一页演示流程关系。四个步骤用箭头形表达先后，序号用圆形节点标注。", BLACK, KAI)], S20,
         anchor=MSO_ANCHOR.MIDDLE, spacing=1.2)
    steps = [("1", "抽取材料", "从材料与来源文件取数"),
             ("2", "写一页纸大纲", "每页一句核心判断"),
             ("3", "生成页面", "骨架与形状按关系选"),
             ("4", "机器检查", "结构、配色、预览三道")]
    left = 0.75
    for number, title, detail in steps:
        rect(slide, left, 3.15, 2.55, 1.30, fill=WHITE, line=MID, shape=MSO_SHAPE.CHEVRON)
        text(slide, left + 0.40, 3.32, 1.95, 0.44, [(title, NAVY, HEI)], S18, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        text(slide, left + 0.34, 3.76, 2.05, 0.66, [(detail, BLACK, SONG)], S16, PP_ALIGN.CENTER,
             MSO_ANCHOR.MIDDLE, spacing=1.15)
        rect(slide, left + 0.24, 2.50, 0.62, 0.62, fill=NAVY, shape=MSO_SHAPE.OVAL)
        text(slide, left + 0.24, 2.50, 0.62, 0.62, [(number, WHITE, HEI)], S16, PP_ALIGN.CENTER,
             MSO_ANCHOR.MIDDLE)
        left += 3.15
    rect(slide, 0.55, 5.30, 12.96, 1.10, fill=BAND)
    rect(slide, 0.55, 5.30, 0.045, 1.10, fill=NAVY)
    text(slide, 0.76, 5.40, 12.54, 0.90,
         [("结论占位：结构变化先改大纲再改页面，这是返工最少的一种改法。", NAVY, HEI)], S18,
         anchor=MSO_ANCHOR.MIDDLE, spacing=1.2)
    slide.notes_slide.notes_text_frame.text = (
        "【流程页 · 备注】本页演示流程关系：箭头形表达先后顺序，圆形序号标记步骤，上导语下结论形成闭环。"
        "讲的时候按步骤顺序走，每一步只说一件事，最后用结论句收束。如果步骤超过五步，先考虑拆成两页，"
        "不要靠缩小字号硬塞。" * 2)

    # ── 第 6 页：时间线 ──
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "06", "示例页：时间关系用一条轴表达",
           "时间线用一条主轴加节点，年份在上、说明在下；不排成一列方块。")
    rect(slide, 0.55, 1.82, 12.96, 0.72, fill=BAND)
    text(slide, 0.76, 1.92, 12.54, 0.52,
         [("导语占位：这一页演示时间线骨架，主轴贯穿全宽，四个节点分别落在轴上。", BLACK, KAI)], S20,
         anchor=MSO_ANCHOR.MIDDLE, spacing=1.2)
    rect(slide, 0.55, 2.68, 12.96, 2.40, fill=BAND)
    rect(slide, 0.85, 3.72, 12.30, 0.045, fill=NAVY)
    nodes = [("示例年 1", "起点：材料与目标对齐"), ("示例年 2", "阶段：结构与方法定稿"),
             ("示例年 3", "推进：数据与证据补齐"), ("示例年 4", "收束：结论与口径复核")]
    left = 1.05
    for index, (year, detail) in enumerate(nodes):
        rect(slide, left + 0.34, 3.46, 0.52, 0.52, fill=NAVY, shape=MSO_SHAPE.OVAL)
        text(slide, left + 0.34, 3.46, 0.52, 0.52, [(str(index + 1), WHITE, HEI)], S16,
             PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        text(slide, left, 2.96, 1.90, 0.44, [(year, NAVY, HEI)], S18, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        text(slide, left, 4.08, 1.90, 0.80, [(detail, BLACK, SONG)], S16, PP_ALIGN.CENTER,
             MSO_ANCHOR.TOP, spacing=1.15)
        left += 3.08
    rect(slide, 0.55, 5.30, 12.96, 1.10, fill=BAND)
    rect(slide, 0.55, 5.30, 0.045, 1.10, fill=SUPPORT_TEXT)
    text(slide, 0.76, 5.40, 12.54, 0.90,
         [("结论占位：时间线只标真正改变方向的事件，不把日常进度也排上去。", NAVY, HEI)], S18,
         anchor=MSO_ANCHOR.MIDDLE, spacing=1.2)
    slide.notes_slide.notes_text_frame.text = (
        "【时间线页 · 备注】时间关系用一条主轴表达，节点在上、说明在下，视觉上能看出先后。"
        "讲的时候只强调真正改变方向的两三个事件，其余可以一句带过。时间线不要用来堆日常进度，"
        "那样等于把待办清单搬上大屏。" * 2)

    # ── 第 7 页：层级 ──
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "07", "示例页：层级关系用宽度变化表达",
           "层级用逐层加宽的条带表达，不用三个同样大的方块并排。")
    layers = [("上层：总目标", "一句话说清这套汇报要证明什么", 6.50, NAVY),
              ("中层：三条主线", "把证据按主线分组，每条只讲一个机制", 8.60, MID),
              ("下层：具体实证", "指标、数值、口径逐条落到同一句里", 10.70, SUPPORT_TEXT)]
    top = 2.10
    for title, detail, width, fill in layers:
        left = 0.55 + (12.96 - width) / 2.0
        rect(slide, left, top, width, 1.05, fill=fill, line=LINE)
        text(slide, left + 0.30, top + 0.12, width - 0.60, 0.44, [(title, WHITE, HEI)], S18,
             anchor=MSO_ANCHOR.MIDDLE)
        text(slide, left + 0.30, top + 0.54, width - 0.60, 0.42, [(detail, WHITE, SONG)], S16,
             anchor=MSO_ANCHOR.MIDDLE)
        top += 1.16
    text(slide, 0.55, 5.72, 12.96, 0.62,
         [("判断占位：层级越往上越概括，越往下越具体；同级信息必须同宽，不靠字号区分层级。", BLACK, SONG)], S18,
         anchor=MSO_ANCHOR.MIDDLE, spacing=1.2)
    slide.notes_slide.notes_text_frame.text = (
        "【层级页 · 备注】层级关系用宽度变化表达：越往上越概括，越往下越具体。讲的时候先给上层结论，"
        "再往下拆，最后落到具体实证。同级信息保持同宽同高，不要用字号大小暗示层级，"
        "否则读者会先去猜哪个更重要。" * 2)

    # ── 第 8 页：循环 / 中心辐射 ──
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "08", "示例页：循环关系用中心与节点表达",
           "闭环、迭代、聚合关系用中心圆加周边节点，不用箭头回环堆满页面。")
    rect(slide, 4.85, 3.10, 3.20, 3.20, fill=NAVY, shape=MSO_SHAPE.OVAL)
    text(slide, 5.15, 3.92, 2.60, 0.72, [("中心主题", WHITE, HEI)], S20, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    text(slide, 5.15, 4.66, 2.60, 0.50, [("一句话定义", WHITE, SONG)], S16, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    nodes = [("节点一", 1.35, 2.70), ("节点二", 10.05, 2.70), ("节点三", 1.35, 4.70), ("节点四", 10.05, 4.70)]
    for label, left, top in nodes:
        rect(slide, left, top, 1.95, 1.20, fill=SUPPORT_FILL, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        text(slide, left + 0.12, top + 0.16, 1.71, 0.88,
             [(label, SUPPORT[0], HEI), ("\n说明占位", BLACK, SONG)], S16, PP_ALIGN.CENTER,
             MSO_ANCHOR.MIDDLE, spacing=1.15)
    text(slide, 0.55, 6.50, 12.96, 0.58,
         [("判断占位：闭环的关键是回到起点时标准变了，而不是把四步原样转一圈。", BLACK, SONG)], S18,
         anchor=MSO_ANCHOR.MIDDLE, spacing=1.2)
    slide.notes_slide.notes_text_frame.text = (
        "【循环页 · 备注】闭环、迭代、聚合关系用中心圆加周边节点表达。讲的时候先给中心概念的准确定义，"
        "再说明每个节点做什么，最后点出回路的价值：回到起点时标准或能力已经变化。"
        "如果四个节点之间其实是先后关系，应该改成流程页或时间线，而不是画成闭环。" * 2)

    # ── 第 9 页：纯文字观点页 ──
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    text(slide, 0.55, 0.24, 12.96, 0.34, [("示例 · 纯文字观点页", BLACK, SONG)], S16, PP_ALIGN.RIGHT)
    rect(slide, 0.55, 1.90, 0.045, 2.60, fill=AMBER_TEXT)
    text(slide, 0.95, 1.95, 12.20, 1.40, [("宁可没有，不能乱加。", NAVY, HEI)], S36, spacing=1.1)
    text(slide, 0.95, 3.55, 12.20, 1.00,
         [("纯文字页允许存在，但必须有字号层级、留白和对齐结构，不能是一大段文字居中堆着。"
           "这一页用 36pt 主句、20pt 次级句和 16pt 小字建立三级层次，靠留白承担视觉任务。", BLACK, KAI)], S20,
         spacing=1.25)
    text(slide, 0.95, 4.80, 12.20, 0.44,
         [("整套演示里纯文字页建议不超过两成；只有核心观点、引用、转折和结尾适合这样处理。", MID, SONG)], S16)
    slide.notes_slide.notes_text_frame.text = (
        "【纯文字页 · 备注】这一页演示纯文字页的合规做法：允许不加图形和配图，但必须有字号层级、留白和对齐。"
        "讲的时候放慢语速，把这句判断说完整，不要急着翻页。整套演示里纯文字页占比建议不超过两成，"
        "超出就说明多数页面没有把关系画出来。" * 2)


def build_negative(prs, small_image):
    """反例：故意违反规则，用于验证检查脚本能发现问题。"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    text(slide, 0.55, 0.60, 12.96, 0.50, [("反例一：字号、配色与字体同时违规", "FF0000", "微软雅黑")], 12)
    text(slide, 0.55, 1.30, 12.96, 0.60,
         [("这一行是 12pt 的 FF0000 红字，用来验证字号、字阶与方案外颜色检查。", "FF0000", "微软雅黑")], 12)
    text(slide, 0.75, 2.20, 12.56, 0.60,
         [("这一行是白底上的浅蓝字，对比度不达标。", LIGHT, SONG)], S18)
    slide.notes_slide.notes_text_frame.text = "反例：本页故意违规。"

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    text(slide, 0.55, 0.60, 12.96, 0.50, [("反例二：字形缺失与文本溢出", NAVY, HEI)], S30)
    text(slide, 0.55, 1.30, 12.96, 0.60,
         [("矩阵转置写成 Vᵀ，奇异值写成 sigma₁，上标平方写成 x²，这些字符超出 GBK 范围。", BLACK, SONG)], S16)
    rect(slide, 0.55, 2.10, 3.00, 0.50, fill=BAND)
    text(slide, 0.60, 2.14, 2.90, 0.42,
         [("这个框高度不足，文字一定会溢出。", BLACK, SONG)], S16)
    slide.notes_slide.notes_text_frame.text = "反例：本页故意违规。"

    for index in range(3):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        text(slide, 0.55, 0.60, 12.96, 0.50, [(f"反例三之{index + 1}：连续三页同一骨架", NAVY, HEI)], S30)
        text(slide, 0.55, 1.40, 12.96, 1.00,
             [("这一页只有一段文字，没有图表、图片或图形结构，字号也没有层级，"
               "用来验证骨架重复与纯文字页比例检查。", BLACK, SONG)], S16)
        slide.notes_slide.notes_text_frame.text = "反例：本页故意违规。"

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    text(slide, 0.55, 0.60, 12.96, 0.50, [("反例四：低分辨率图片", NAVY, HEI)], S30)
    slide.shapes.add_picture(small_image, Inches(0.55), Inches(1.40), Inches(6.00), Inches(3.40))
    text(slide, 6.90, 1.40, 6.60, 1.00,
         [("这张图只有 160×90 像素，却被放大到 6 英寸宽，实际约 27 dpi，用来验证配图分辨率检查。",
           BLACK, SONG)], S16)
    slide.notes_slide.notes_text_frame.text = "反例：本页故意违规。"


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--variant", choices=("positive", "negative"), default="positive")
    parser.add_argument("--out", default=None)
    args = parser.parse_args()

    os.makedirs(ASSETS, exist_ok=True)
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    if args.variant == "positive":
        out = args.out or os.path.join(ASSETS, "example-slide.pptx")
        placeholder = os.path.join(ASSETS, "placeholder-photo.png")
        if not os.path.exists(placeholder):
            write_png(placeholder, 1600, 900, gradient_columns(1600))
        build_positive(prs, placeholder)
    else:
        out = args.out or os.path.join(os.path.dirname(os.path.abspath(__file__)), "example-negative.pptx")
        small = os.path.join(os.path.dirname(os.path.abspath(out)), "placeholder-small.png")
        if not os.path.exists(small):
            write_png(small, 160, 90, gradient_columns(160, end=CONTRAST[3]), red_band=True)
        build_negative(prs, small)

    out_dir = os.path.dirname(os.path.abspath(out))
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    prs.save(out)
    print("saved", os.path.normpath(out))


if __name__ == "__main__":
    main()