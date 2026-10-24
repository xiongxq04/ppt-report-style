# -*- coding: utf-8 -*-
"""校验 pptx 是否符合汇报规范：字阶、配色方案与对比度、中文字体（ea 槽）与西文字体（latin 槽）双白名单。

用法:
  python verify_ppt_style.py <file.pptx> [更多文件 ...] [选项]
选项:
  --mode balanced|strict  平衡档默认，严格档把继承项也判为错误
  --accent HEX            主题色，默认 283E4E
  --contrast-color HEX    对比色基色，默认用内置暖琥珀系列
  --support HEX           辅助色基色，默认用内置灰青绿系列
  --allow-colors HEX      额外允许的颜色（逗号分隔）
  --min-size N            最小字号(pt)，默认 16
  --fonts "a,b,c"         允许的中文字体（a:ea），默认 宋体,楷体,黑体
  --latin-fonts "x,y"     允许的西文字体（a:latin），默认 Arial,Times New Roman
  --sizes "a,b,..."       允许的字阶档位，默认 16,18,20,30,36,44
  --max-steps N           每页最多几档字号，默认 5
  --allow-cjk-latin       容忍西文槽填中文字体（只用于跑旧稿）
  --max-per-rule N        同页同规则的展示条数，0 表示不折叠
退出码 0 表示没有错误。

为什么分两槽：python-pptx 的 run.font.name 写的是 a:latin，中文实际取 a:ea。
只写 font.name=宋体 → 英文数字用宋体的西文字形，字距松、字形丑，本项检查也会漏掉。
配色规则：黑白灰（中性色）不占配额；主题色、对比色、辅助色各是一个色阶系列。详见 references/palette-system.md。
"""
import argparse
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.stdout.reconfigure(encoding="utf-8")

from palette_utils import (ACCENT_AREA_MAX, ACCENT_AREA_WARN, ACCENT_DEFAULT, FAMILY_LABEL,
                           MAX_HUE_GROUPS, SUPPORT_AREA_MAX, SUPPORT_AREA_WARN,
                           composite, contrast_ratio, family_of, hue_group, is_neutral,
                           normalize_hex, required_ratio)

try:
    from pptx import Presentation
    from pptx.enum.dml import MSO_FILL
    from pptx.oxml.ns import qn
except ImportError:
    print("需要 python-pptx：pip install python-pptx")
    sys.exit(2)

EMU_IN = 914400.0
SEVERITY_ORDER = {"ERROR": 0, "WARNING": 1, "SUGGESTION": 2, "例外": 3}
CJK_FACES = {"宋体", "楷体", "黑体", "SimSun", "KaiTi", "SimHei", "MS Mincho", "Microsoft YaHei"}
BASE = {"000000", "FFFFFF"}


def has_cjk(text):
    return any("\u3000" <= ch <= "\u9fff" or "\uff00" <= ch <= "\uffef" for ch in text)


def has_latin(text):
    return any(ch.isalnum() and not has_cjk(ch) for ch in text)


def slots(run):
    """返回 (latin 槽, ea 槽)；未显式设置的返回空串。"""
    rpr = run._r.find(qn("a:rPr"))
    if rpr is None:
        return (run.font.name or "", "")
    latin = rpr.find(qn("a:latin"))
    ea = rpr.find(qn("a:ea"))
    return ((latin.get("typeface") if latin is not None else None) or (run.font.name or ""),
            ea.get("typeface") if ea is not None else "")


def explicit_rgb(color_format):
    try:
        return normalize_hex(str(color_format.rgb), None)
    except Exception:
        return None


def fill_alpha(element):
    try:
        node = element.spPr.find(qn("a:solidFill"))
        if node is None:
            return 1.0
        srgb = node.find(qn("a:srgbClr"))
        if srgb is None:
            return 1.0
        alpha = srgb.find(qn("a:alpha"))
        if alpha is not None and alpha.get("val"):
            return int(alpha.get("val")) / 100000.0
    except Exception:
        pass
    return 1.0


def solid_fill(owner):
    try:
        fill = owner.fill
        if fill.type != MSO_FILL.SOLID:
            return None
        value = explicit_rgb(fill.fore_color)
    except Exception:
        return None
    if value is None:
        return None
    try:
        return value, fill_alpha(owner._element)
    except Exception:
        return value, 1.0


def shape_rect(shape):
    try:
        return (shape.left / EMU_IN, shape.top / EMU_IN,
                shape.width / EMU_IN, shape.height / EMU_IN)
    except Exception:
        return None


def slide_background(slide, default="FFFFFF"):
    layout = getattr(slide, "slide_layout", None)
    master = getattr(layout, "slide_master", None)
    for holder in (slide, layout, master):
        if holder is None:
            continue
        try:
            fill = holder.background.fill
            if fill.type == MSO_FILL.SOLID:
                value = explicit_rgb(fill.fore_color)
                if value:
                    return value
        except Exception:
            continue
    return default


def backdrop_for(shape, index, slide_bg, backdrops):
    """文字实际压在什么颜色上：优先自身填充，其次压在下面的最后一个填充形状。"""
    own = solid_fill(shape)
    if own is not None:
        return composite(own[0], slide_bg, own[1])
    rect = shape_rect(shape)
    if rect is None:
        return slide_bg
    center_x = rect[0] + rect[2] / 2.0
    center_y = rect[1] + rect[3] / 2.0
    best = None
    for position, color, alpha, box in backdrops:
        if position >= index:
            continue
        if box[0] <= center_x <= box[0] + box[2] and box[1] <= center_y <= box[1] + box[3]:
            if best is None or position > best[0]:
                best = (position, color, alpha)
    if best is None:
        return slide_bg
    return composite(best[1], slide_bg, best[2])


def check(path, accent=ACCENT_DEFAULT, min_size=16.0, fonts="宋体,楷体,黑体", allow=None,
          mode="balanced", contrast_color=None, support=None, max_per_rule=3,
          latin_fonts="Arial,Times New Roman", sizes="16,18,20,30,36,44", max_steps=5,
          allow_cjk_latin=False):
    ea_ok = {item.strip() for item in fonts.split(",") if item.strip()}
    lat_ok = {item.strip() for item in latin_fonts.split(",") if item.strip()}
    size_ok = {float(item) for item in sizes.split(",") if item.strip()}
    allowed_extra = set()
    for item in (allow or []):
        value = normalize_hex(item, None)
        if value:
            allowed_extra.add(value)

    findings = []
    prs = Presentation(path)
    slide_area = (prs.slide_width / EMU_IN) * (prs.slide_height / EMU_IN)

    def add(slide_no, severity, rule, kind, detail, extra=""):
        findings.append((severity, slide_no, rule, kind, detail, extra))

    def soft(severity):
        return severity if mode == "strict" else "WARNING"

    def judge_color(slide_no, rule, kind, value, detail):
        family = family_of(value, accent=accent, contrast=contrast_color, support=support)
        if family == "unknown" and value not in allowed_extra and value not in BASE:
            add(slide_no, "ERROR", rule, kind, detail,
                f"#{value} 不在配色方案内（{FAMILY_LABEL['unknown']}）")
        return family

    def check_runs(slide_no, runs, where, background, touch_steps=None, in_table=False):
        for run in runs:
            text = run.text
            if not text.strip():
                continue
            size = run.font.size.pt if run.font.size else None
            if size is None:
                add(slide_no, soft("ERROR"), "TYP.INHERIT", "字号继承", f"{where}{text[:20]}",
                    "未显式设置字号")
            elif size < min_size:
                add(slide_no, "ERROR", "TYP.MINSIZE", "字号", f"{where}{text[:20]}",
                    f"{size}pt < {min_size}pt")
            elif size_ok and size not in size_ok:
                add(slide_no, "ERROR", "TYP.LADDER", "字号不在字阶内", f"{where}{text[:20]}",
                    f"{size}pt，字阶={sorted(size_ok)}")
            if size is not None and touch_steps is not None:
                touch_steps.add(size)

            color = explicit_rgb(run.font.color)
            if color is None:
                add(slide_no, soft("ERROR"), "PAL.INHERIT", "颜色继承", f"{where}{text[:20]}",
                    "主题色或继承色，无法判定对比度")
            else:
                judge_color(slide_no, "PAL.COLOR", "文字颜色", color, f"{where}{text[:20]}")
                ratio = contrast_ratio(color, background)
                need = required_ratio(size, bool(run.font.bold))
                if ratio < need:
                    add(slide_no, "ERROR", "PAL.CONTRAST", "文字对比度",
                        f"{where}#{color} on #{background}",
                        f"{ratio:.2f}:1 < {need:.1f}:1，"
                        f"{'正文口径' if need == 4.5 else '大字口径'}（{size or '?'}pt）")

            latin, ea = slots(run)
            if has_cjk(text):
                if not ea:
                    add(slide_no, soft("ERROR"), "FONT.EA_MISSING", "中文字体未设置",
                        f"{where}{text[:16]}", "a:ea 空，将继承主题字体")
                elif ea_ok and ea not in ea_ok:
                    add(slide_no, "ERROR", "FONT.EA", "中文字体", f"{where}{text[:16]}", ea)
                if ea == "黑体" and run.font.bold:
                    add(slide_no, "ERROR", "FONT.EA_BOLD", "黑体加粗", f"{where}{text[:16]}",
                        "合成粗会糊边；强调改用字号或主题色")
                if in_table and ea == "楷体":
                    add(slide_no, "ERROR", "FONT.TABLE_KAI", "楷体入表格", f"{where}{text[:16]}",
                        "表格只用宋体 / 黑体")
            if has_latin(text):
                if not latin:
                    add(slide_no, soft("ERROR"), "FONT.LATIN_MISSING", "西文字体未设置",
                        f"{where}{text[:16]}", "a:latin 空")
                elif latin in CJK_FACES:
                    if not allow_cjk_latin:
                        add(slide_no, "ERROR", "FONT.LATIN_CJK", "西文用中文字体", f"{where}{text[:16]}",
                            f"latin={latin}（英文将用中文字形渲染）")
                elif lat_ok and latin not in lat_ok:
                    add(slide_no, "ERROR", "FONT.LATIN", "西文字体", f"{where}{text[:16]}", latin)

    for slide_no, slide in enumerate(prs.slides, 1):
        slide_bg = slide_background(slide)
        shapes = list(slide.shapes)
        backdrops = []
        for position, shape in enumerate(shapes):
            info = solid_fill(shape)
            rect = shape_rect(shape)
            if info is not None and rect is not None:
                backdrops.append((position, info[0], info[1], rect))

        used_colors = set()
        areas = {"theme": 0.0, "contrast": 0.0, "support": 0.0}
        steps = set()

        for index, shape in enumerate(shapes):
            shape_bg = backdrop_for(shape, index, slide_bg, backdrops)
            info = solid_fill(shape)
            if info is not None:
                value, alpha = info
                used_colors.add(composite(value, slide_bg, alpha))
                family = judge_color(slide_no, "PAL.COLOR", "形状填充", value, str(shape.shape_type))
                if family in areas:
                    rect = shape_rect(shape)
                    if rect is not None:
                        areas[family] += rect[2] * rect[3]
            try:
                if shape.line.fill.type == MSO_FILL.SOLID:
                    stroke = explicit_rgb(shape.line.color)
                    if stroke:
                        used_colors.add(stroke)
                        judge_color(slide_no, "PAL.COLOR", "形状描边", stroke, str(shape.shape_type))
            except Exception:
                pass

            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                runs = [run for paragraph in shape.text_frame.paragraphs for run in paragraph.runs]
                check_runs(slide_no, runs, "", shape_bg, steps)

            if getattr(shape, "has_table", False):
                for row_no, row in enumerate(shape.table.rows, 1):
                    for col_no, cell in enumerate(row.cells, 1):
                        cell_fill = solid_fill(cell)
                        cell_bg = shape_bg if cell_fill is None else composite(cell_fill[0], shape_bg, cell_fill[1])
                        if cell_fill is not None:
                            used_colors.add(cell_bg)
                            family = judge_color(slide_no, "PAL.COLOR", "单元格填充", cell_fill[0],
                                                 f"表格{row_no},{col_no}")
                            if family in areas:
                                try:
                                    areas[family] += (cell.width / EMU_IN) * (cell.height / EMU_IN)
                                except Exception:
                                    pass
                        if cell.text_frame.text.strip():
                            runs = [run for paragraph in cell.text_frame.paragraphs for run in paragraph.runs]
                            check_runs(slide_no, runs, f"表格{row_no},{col_no}: ", cell_bg, steps, in_table=True)

            if getattr(shape, "has_chart", False):
                chart = shape.chart
                seen = set()
                for tag in ("a:defRPr", "a:rPr"):
                    for node in chart._chartSpace.iter(qn(tag)):
                        key = (tag, id(node))
                        if key in seen:
                            continue
                        seen.add(key)
                        size = node.get("sz")
                        if size and int(size) < min_size * 100:
                            add(slide_no, "ERROR", "TYP.MINSIZE", "图表字号", tag,
                                f"{int(size) / 100}pt < {min_size}pt")
                        elif not size:
                            add(slide_no, soft("ERROR"), "TYP.INHERIT", "图表字号继承", tag,
                                "未显式设置字号")
                        fill = node.find(qn("a:solidFill"))
                        if fill is not None:
                            srgb = fill.find(qn("a:srgbClr"))
                            if srgb is not None:
                                color = normalize_hex(srgb.get("val"), None)
                                if color:
                                    used_colors.add(color)
                                    judge_color(slide_no, "PAL.COLOR", "图表颜色", color, tag)
                                    ratio = contrast_ratio(color, slide_bg)
                                    if ratio < required_ratio(float(size) / 100 if size else None):
                                        add(slide_no, "ERROR", "PAL.CONTRAST", "图表文字对比度",
                                            f"#{color} on #{slide_bg}", f"{ratio:.2f}:1")

        if len(steps) > max_steps:
            add(slide_no, "ERROR", "TYP.STEPS", "字号层级过多",
                f"第{slide_no}页用了 {len(steps)} 档字号", f"上限 {max_steps} 档：{sorted(steps)}")

        hues = sorted({hue_group(color) for color in used_colors if not is_neutral(color)})
        if len(hues) > MAX_HUE_GROUPS:
            add(slide_no, soft("ERROR"), "PAL.HUE", "色相组数",
                f"第{slide_no}页使用了 {len(hues)} 组色相",
                f"超出单页 {MAX_HUE_GROUPS} 组上限；黑白灰不计入")

        if slide_area > 0:
            for family, warn, ceiling, label in (
                ("contrast", ACCENT_AREA_WARN, ACCENT_AREA_MAX, "对比色"),
                ("support", SUPPORT_AREA_WARN, SUPPORT_AREA_MAX, "辅助色"),
            ):
                share = areas[family] / slide_area
                if share > ceiling:
                    add(slide_no, "ERROR", f"PAL.AREA_{family.upper()}", f"{label}面积",
                        f"第{slide_no}页 {label}约占 {share:.0%}", f"超过上限 {ceiling:.0%}")
                elif share > warn:
                    add(slide_no, "WARNING", f"PAL.AREA_{family.upper()}", f"{label}面积",
                        f"第{slide_no}页 {label}约占 {share:.0%}", f"建议不超过 {warn:.0%}")

    if max_per_rule:
        grouped = {}
        for item in findings:
            grouped.setdefault((item[1], item[2]), []).append(item)
        compressed = []
        for (slide_no, rule), items in grouped.items():
            if len(items) <= max_per_rule:
                compressed.extend(items)
                continue
            compressed.extend(items[:max_per_rule])
            severity = min(items, key=lambda item: SEVERITY_ORDER.get(item[0], 3))[0]
            compressed.append((severity, slide_no, rule, "同类问题",
                               f"另有 {len(items) - max_per_rule} 处同类问题",
                               "用 --max-per-rule 0 查看完整清单"))
        findings = compressed
    return findings


def format_findings(findings):
    if not findings:
        return "OK"
    lines = []
    for severity, slide_no, rule, kind, detail, extra in sorted(
            findings, key=lambda item: (item[1], SEVERITY_ORDER.get(item[0], 3), item[2])):
        where = f"第{slide_no}页" if slide_no else "全篇"
        lines.append(f"{severity}  {where}  [{rule}] {kind}: {detail}" + (f" — {extra}" if extra else ""))
    return "\n".join(lines)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("files", nargs="*")
    parser.add_argument("--accent", default=ACCENT_DEFAULT)
    parser.add_argument("--contrast-color", default=None)
    parser.add_argument("--support", default=None)
    parser.add_argument("--min-size", type=float, default=16)
    parser.add_argument("--fonts", default="宋体,楷体,黑体", help="中文字体（a:ea）白名单")
    parser.add_argument("--latin-fonts", default="Arial,Times New Roman", help="西文字体（a:latin）白名单")
    parser.add_argument("--allow-colors", default="")
    parser.add_argument("--allow-cjk-latin", action="store_true", help="容忍西文槽填中文字体（旧稿用）")
    parser.add_argument("--sizes", default="16,18,20,30,36,44", help="允许的字阶档位，逗号分隔")
    parser.add_argument("--max-steps", type=int, default=5, help="每页最多出现几档字号")
    parser.add_argument("--mode", choices=("balanced", "strict"), default="balanced")
    parser.add_argument("--max-per-rule", type=int, default=3,
                        help="同一页同一规则的展示条数，0 表示不折叠")
    args = parser.parse_args()
    if not args.files:
        print(__doc__)
        return 2

    allow = [item.strip() for item in args.allow_colors.split(",") if item.strip()]
    exit_code = 0
    for path in args.files:
        if not os.path.exists(path):
            print("文件不存在:", path)
            exit_code = 1
            continue
        findings = check(path, args.accent, args.min_size, args.fonts, allow, args.mode,
                         args.contrast_color, args.support, args.max_per_rule,
                         args.latin_fonts, args.sizes, args.max_steps, args.allow_cjk_latin)
        errors = [item for item in findings if item[0] == "ERROR"]
        print(("FAIL" if errors else "OK  "), os.path.basename(path), f"({len(findings)} 项)")
        if findings:
            print(format_findings(findings))
        if errors:
            exit_code = 1
    return exit_code


if __name__ == "__main__":
    raise SystemExit(main())